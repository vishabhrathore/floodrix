import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    
    for name in wb.sheetnames:
        ws = wb[name]
        for r in range(1, 150):
            for c in range(1, 150):
                val = ws.cell(r, c).value
                if val is not None and isinstance(val, str):
                    if any(x in val.lower() for x in ["peak discharge", "time to peak", "time to rise", "base width", "w50", "w75", "wr50", "wr75"]):
                        print(f"Sheet '{name}', Cell {openpyxl.utils.get_column_letter(c)}{r}: {val} = {ws.cell(r, c+1).value}")

if __name__ == "__main__":
    main()
