import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["24hr SPS"]
    
    print("=== Formulas ===")
    for r in range(4, 12):
        row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(r, c).value for c in range(7, 10)}
        print(f"Row {r:02d}: {row_vals}")
        
    wb_data = openpyxl.load_workbook(path, data_only=True)
    print("=== Values ===")
    for r in range(4, 12):
        row_vals = {openpyxl.utils.get_column_letter(c): wb_data["24hr SPS"].cell(r, c).value for c in range(7, 10)}
        print(f"Row {r:02d}: {row_vals}")

if __name__ == "__main__":
    main()
