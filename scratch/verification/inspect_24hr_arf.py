import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    print("24hr SPS ARF value:", wb["24hr SPS"]["E11"].value) # let's print around E11/I9
    
    # search for ARF in sheet 24hr SPS
    ws = wb["24hr SPS"]
    for r in range(1, 25):
        for c in range(1, 15):
            val = ws.cell(r, c).value
            if val and "reduction" in str(val).lower():
                print(f"Cell {openpyxl.utils.get_column_letter(c)}{r}: {val}")
                # print next col
                print(f"Next cell: {ws.cell(r, c+1).value}")

if __name__ == "__main__":
    main()
