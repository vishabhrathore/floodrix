import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["24hr SPS"]
    
    max_val = -1
    max_row = -1
    for r in range(53, 200):
        val = ws.cell(r, 27).value
        if val is not None and val > max_val:
            max_val = val
            max_row = r
            
    print(f"Max DRH (Col AA) is {max_val} at row {max_row}")
    print(f"Base flow at row {max_row} is {ws.cell(max_row, 28).value}")
    print(f"PMF at row {max_row} is {ws.cell(max_row, 29).value}")

if __name__ == "__main__":
    main()
