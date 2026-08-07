import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["24hr SPS"]
    
    print("=== Row 52 ===")
    for c in range(1, 53):
        val = ws.cell(52, c).value
        if val is not None:
            print(f"Col {openpyxl.utils.get_column_letter(c)}: {val}")
            
    print("=== Row 53 ===")
    for c in range(1, 53):
        val = ws.cell(53, c).value
        if val is not None:
            print(f"Col {openpyxl.utils.get_column_letter(c)}: {val}")

if __name__ == "__main__":
    main()
