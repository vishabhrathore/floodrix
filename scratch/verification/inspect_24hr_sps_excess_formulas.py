import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["24hr SPS"]
    
    print("=== 24hr SPS excess formulas (rows 35 to 46, cols N to Q) ===")
    for row in range(35, 47):
        row_vals = {openpyxl.utils.get_column_letter(col): ws.cell(row, col).value for col in range(14, 18)}
        print(f"Row {row:02d}: {row_vals}")

if __name__ == "__main__":
    main()
