import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["24hr SPS"]
    
    print("=== Row 53 formulas ===")
    print({openpyxl.utils.get_column_letter(c): ws.cell(53, c).value for c in range(12, 19)})
    
    print("=== Row 94 formulas ===")
    print({openpyxl.utils.get_column_letter(c): ws.cell(94, c).value for c in range(12, 19)})
    
    wb_data = openpyxl.load_workbook(path, data_only=True)
    print("=== Row 94 values ===")
    print({openpyxl.utils.get_column_letter(c): wb_data["24hr SPS"].cell(94, c).value for c in range(12, 19)})

if __name__ == "__main__":
    main()
