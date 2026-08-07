import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["48hr SPS "]
    
    print("=== 48hr SPS R35:U35 array formulas ===")
    for col in range(18, 22):
        cell = ws.cell(35, col)
        val = cell.value
        # get text attribute
        print(f"Col {openpyxl.utils.get_column_letter(col)}: {val.text if hasattr(val, 'text') else str(val)}")

if __name__ == "__main__":
    main()
