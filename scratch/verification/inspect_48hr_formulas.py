import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["48hr SPS "]
    
    print("=== 48hr SPS formulas ===")
    for c in range(12, 24):
        col_let = openpyxl.utils.get_column_letter(c)
        print(f"Col {col_let} row 35: {ws.cell(35, c).value}")
        print(f"Col {col_let} row 36: {ws.cell(36, c).value}")

if __name__ == "__main__":
    main()
