import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS I36:N36 formulas ===")
    for col in range(9, 15):
        cell = ws.cell(36, col)
        print(f"Col {openpyxl.utils.get_column_letter(col)}: ref={getattr(cell, 'ref', None)}, val={cell.value}")

if __name__ == "__main__":
    main()
