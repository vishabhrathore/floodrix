import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS P36:U36 formulas ===")
    for col in range(16, 22):
        cell = ws.cell(36, col)
        print(f"Col {openpyxl.utils.get_column_letter(col)}: {cell.value}")

if __name__ == "__main__":
    main()
