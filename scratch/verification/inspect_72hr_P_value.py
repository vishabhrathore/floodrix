import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS P35:U48 values ===")
    for r in range(35, 49):
        row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(r, c).value for c in range(16, 22)}
        print(f"Row {r:02d}: {row_vals}")

if __name__ == "__main__":
    main()
