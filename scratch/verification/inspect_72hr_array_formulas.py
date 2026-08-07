import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS V36:AA36 formulas ===")
    for col in range(22, 28):
        cell = ws.cell(36, col)
        val = cell.value
        # Check attributes of val
        import inspect
        attrs = {name: getattr(val, name, None) for name, _ in inspect.getmembers(val) if not name.startswith("_")}
        print(f"Col {openpyxl.utils.get_column_letter(col)}: {attrs}")

if __name__ == "__main__":
    main()
