import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["72hr SPS  "]
    
    print("=== Search for text in 72hr SPS sheet ===")
    for r in range(1, 100):
        for c in range(1, 20):
            val = ws.cell(r, c).value
            if val and any(word in str(val).lower() for word in ["best", "order", "max"]):
                print(f"Cell {openpyxl.utils.get_column_letter(c)}{r}: {val}")

if __name__ == "__main__":
    main()
