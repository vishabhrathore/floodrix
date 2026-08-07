import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS I35:O48 formulas ===")
    for r in range(35, 49):
        row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(r, c).value for c in range(9, 16)}
        print(f"Row {r:02d}: {row_vals}")

if __name__ == "__main__":
    main()
