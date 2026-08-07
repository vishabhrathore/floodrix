import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS B53:B65 values ===")
    for r in range(53, 66):
        print(f"Row {r}: B={ws.cell(r, 2).value}")
        
    wb_form = openpyxl.load_workbook(path, data_only=False)
    print("=== 72hr SPS B53:B65 formulas ===")
    for r in range(53, 66):
        print(f"Row {r}: B={wb_form['72hr SPS  '].cell(r, 2).value}")

if __name__ == "__main__":
    main()
