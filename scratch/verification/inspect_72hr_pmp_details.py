import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb_val = openpyxl.load_workbook(path, data_only=True)
    wb_form = openpyxl.load_workbook(path, data_only=False)
    
    ws_val = wb_val["72hr SPS  "]
    ws_form = wb_form["72hr SPS  "]
    
    for r in range(1, 15):
        val = ws_val.cell(r, 9).value
        form = ws_form.cell(r, 9).value
        label = ws_val.cell(r, 7).value
        print(f"Row {r:02d}: label={label}, val={val}, formula={form}")

if __name__ == "__main__":
    main()
