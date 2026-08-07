import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS Row 52 C:H formulas ===")
    for col in range(3, 9):
        print(f"Col {openpyxl.utils.get_column_letter(col)}: {ws.cell(52, col).value}")

if __name__ == "__main__":
    main()
