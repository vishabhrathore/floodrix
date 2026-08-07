import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["72hr SPS  "]
    
    print("=== Row 53 formulas ===")
    print({openpyxl.utils.get_column_letter(c): ws.cell(53, c).value for c in range(70, 80)})
    
    print("=== Row 94 formulas ===")
    print({openpyxl.utils.get_column_letter(c): ws.cell(94, c).value for c in range(70, 80)})

if __name__ == "__main__":
    main()
