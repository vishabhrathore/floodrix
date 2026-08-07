import openpyxl

def check_max(sheet_name, drh_col_idx, pmf_col_idx, start_row, end_row):
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    
    max_drh = -1
    max_drh_row = -1
    for r in range(start_row, end_row + 1):
        val = ws.cell(r, drh_col_idx).value
        if val is not None and isinstance(val, (int, float)) and val > max_drh:
            max_drh = val
            max_drh_row = r
            
    max_pmf = -1
    max_pmf_row = -1
    for r in range(start_row, end_row + 1):
        val = ws.cell(r, pmf_col_idx).value
        if val is not None and isinstance(val, (int, float)) and val > max_pmf:
            max_pmf = val
            max_pmf_row = r
            
    print(f"=== {sheet_name} ===")
    print(f"Max DRH (Col {openpyxl.utils.get_column_letter(drh_col_idx)}) is {max_drh} at row {max_drh_row}")
    print(f"Max PMF (Col {openpyxl.utils.get_column_letter(pmf_col_idx)}) is {max_pmf} at row {max_pmf_row}")

def main():
    # 24hr: DRH is AA (27), PMF is AC (29)
    check_max("24hr SPS", 27, 29, 53, 193)
    # 48hr: DRH is AY (51), PMF is BA (53)
    check_max("48hr SPS ", 51, 53, 53, 193)
    # 72hr: DRH is BW (75), PMF is BY (77)
    check_max("72hr SPS  ", 75, 77, 53, 193)

if __name__ == "__main__":
    main()
