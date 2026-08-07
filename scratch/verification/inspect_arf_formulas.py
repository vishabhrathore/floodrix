import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["24hr SPS"]
    
    print("=== 24hr SPS ARF table formulas (rows 30 to 38) ===")
    for r in range(30, 39):
        row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(r, c).value for c in range(12, 16)}
        print(f"Row {r:02d}: {row_vals}")

if __name__ == "__main__":
    main()
