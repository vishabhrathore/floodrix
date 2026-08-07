import openpyxl

def print_bands(sheet_name):
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    
    print(f"=== {sheet_name} Bands ===")
    # The bands are in rows 4 to 12 or similar
    for r in range(1, 20):
        # print cols A to E
        row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(r, c).value for c in range(1, 6)}
        print(f"Row {r:02d}: {row_vals}")

def main():
    print_bands("48hr SPS ")
    print_bands("72hr SPS  ")

if __name__ == "__main__":
    main()
