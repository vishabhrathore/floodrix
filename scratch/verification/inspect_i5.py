import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    
    for sname in ["24hr SPS", "48hr SPS ", "72hr SPS  "]:
        ws = wb[sname]
        print(f"=== Sheet '{sname}' Cell I5 Formula: {ws['I5'].value}")
        # Let's also print the value in data_only mode
        wb_val = openpyxl.load_workbook(path, data_only=True)
        print(f"=== Sheet '{sname}' Cell I5 Value: {wb_val[sname]['I5'].value}")

if __name__ == "__main__":
    main()
