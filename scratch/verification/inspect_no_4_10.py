import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["24hr SPS"]
    
    print("=== 24hr SPS N4:O10 ===")
    for r in range(4, 11):
        print(f"Row {r}: N={ws.cell(r, 14).value}, O={ws.cell(r, 15).value}")

if __name__ == "__main__":
    main()
