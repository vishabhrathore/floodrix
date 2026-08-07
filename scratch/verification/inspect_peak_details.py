import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    
    # Let's inspect 24hr SPS Row 70
    ws24 = wb["24hr SPS"]
    print("24hr SPS Row 70:")
    row70_vals = {
        "DRH": ws24.cell(70, 27).value,  # AA
        "BaseFlow": ws24.cell(70, 28).value,  # AB
        "PMF": ws24.cell(70, 29).value  # AC
    }
    print(row70_vals)
    
    # Let's inspect 48hr SPS Row 94
    ws48 = wb["48hr SPS "]
    print("48hr SPS Row 94:")
    row94_vals = {
        "DRH": ws48.cell(94, 51).value,  # AY
        "BaseFlow": ws48.cell(94, 52).value,  # AZ
        "PMF": ws48.cell(94, 53).value  # BA
    }
    print(row94_vals)
    
    # Let's inspect 72hr SPS Row 94
    ws72 = wb["72hr SPS  "]
    print("72hr SPS Row 94:")
    row94_72 = {
        "DRH": ws72.cell(94, 75).value,  # BW
        "BaseFlow": ws72.cell(94, 76).value,  # BX
        "PMF": ws72.cell(94, 77).value  # BY
    }
    print(row94_72)

if __name__ == "__main__":
    main()
