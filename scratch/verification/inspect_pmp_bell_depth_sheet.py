import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb_val = openpyxl.load_workbook(path, data_only=True)
    wb_form = openpyxl.load_workbook(path, data_only=False)
    
    ws_val = wb_val["PMP 12hr Bell depth"]
    ws_form = wb_form["PMP 12hr Bell depth"]
    
    print("=== PMP 12hr Bell depth sheet ===")
    for r in range(1, 6):
        form_row = {openpyxl.utils.get_column_letter(c): ws_form.cell(r, c).value for c in range(1, 16)}
        val_row = {openpyxl.utils.get_column_letter(c): ws_val.cell(r, c).value for c in range(1, 16)}
        print(f"Row {r:02d} Formula: {form_row}")
        print(f"Row {r:02d} Value:   {val_row}")

if __name__ == "__main__":
    main()
