import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["PMP 12hr Bell depth"]
    
    print("=== Row 2 Headers ===")
    print([ws.cell(2, c).value for c in range(1, 16)])
    print("=== Row 3 Formulas ===")
    print([ws.cell(3, c).value for c in range(1, 16)])

if __name__ == "__main__":
    main()
