import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    sheet = wb["SUH"]
    
    print("=== SUH sheet profile formulas (rows 20 to 35) ===")
    for row in range(20, 36):
        row_vals = {openpyxl.utils.get_column_letter(col): sheet.cell(row, col).value for col in range(20, 27)}
        print(f"Row {row:02d}: {row_vals}")

if __name__ == "__main__":
    main()
