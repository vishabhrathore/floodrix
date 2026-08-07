import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    sheet = wb["SUH"]
    
    print("=== Row 19 formulas ===")
    row_vals = {openpyxl.utils.get_column_letter(col): sheet.cell(19, col).value for col in range(20, 27)}
    print(row_vals)

if __name__ == "__main__":
    main()
