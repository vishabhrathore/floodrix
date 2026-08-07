import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    
    for sname in ["24hr SPS", "48hr SPS ", "72hr SPS  "]:
        ws = wb[sname]
        print(f"\n================ SHEET: {sname} Row 52 Columns C..AX/AY/BW ================")
        
        # Let's see non-empty values in row 52 starting from column C (col 3)
        cols_val = []
        for col in range(3, 100):
            val = ws.cell(52, col).value
            if val is not None:
                cols_val.append((openpyxl.utils.get_column_letter(col), val))
        print(f"Num columns: {len(cols_val)}")
        print("First 15 cols:", cols_val[:15])
        print("Last 15 cols:", cols_val[-15:])

if __name__ == "__main__":
    main()
