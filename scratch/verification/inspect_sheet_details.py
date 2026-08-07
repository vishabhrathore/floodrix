import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    
    for sname in ["24hr SPS", "48hr SPS ", "72hr SPS  "]:
        ws = wb[sname]
        print(f"\n================ SHEET: {sname} ================")
        
        # Look for "Best order" or "Max" report cells
        for r in range(40, 53):
            row_vals = [ws.cell(r, c).value for c in range(1, 10)]
            if any(row_vals):
                print(f"Row {r}: {row_vals}")
                
        # Let's find the peak convolved value and its row/col
        # For 24hr SPS, sum is in AA, rows 53..200
        # For 48hr SPS, sum is in AY, rows 53..200
        # For 72hr SPS, sum is in BW, rows 53..200
        sum_col = "AA" if sname == "24hr SPS" else ("AY" if "48" in sname else "BW")
        max_val = -1
        max_row = -1
        col_idx = openpyxl.utils.column_index_from_string(sum_col)
        for r in range(53, 500):
            val = ws.cell(r, col_idx).value
            if val is not None and isinstance(val, (int, float)):
                if val > max_val:
                    max_val = val
                    max_row = r
        print(f"Calculated Max in Column {sum_col}: {max_val} at Row {max_row}")

if __name__ == "__main__":
    main()
