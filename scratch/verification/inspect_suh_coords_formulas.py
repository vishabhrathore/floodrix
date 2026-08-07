import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    sheet = wb["SUH"]
    
    print("=== Table 3 (Table of key coordinates) ===")
    for row in range(70, 77):
        print(f"Row {row:02d}: T={sheet.cell(row, 20).value}, U={sheet.cell(row, 21).value}, V={sheet.cell(row, 22).value}")
        
    print("\n=== Table 4 (Interpolated ordinates) ===")
    for row in range(78, 90):
        print(f"Row {row:02d}: T={sheet.cell(row, 20).value}, U={sheet.cell(row, 21).value}, V={sheet.cell(row, 22).value}")

if __name__ == "__main__":
    main()
