import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    sheet = wb["SUH"]
    
    print("=== SUH sheet formulas (cols T to AB, rows 40 to 51) ===")
    for row in range(40, 52):
        row_vals = {openpyxl.utils.get_column_letter(col): sheet.cell(row, col).value for col in range(20, 29)}
        print(f"Row {row:02d}: {row_vals}")

if __name__ == "__main__":
    main()
