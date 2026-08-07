import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["72hr SPS  "]
    
    print("=== 72hr SPS I36:N47 values ===")
    for r in range(36, 48):
        row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(r, c).value for c in range(9, 15)}
        print(f"Row {r:02d}: {row_vals}")

if __name__ == "__main__":
    main()
