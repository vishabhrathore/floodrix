import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["72hr SPS  "]
    
    print("=== Row 35 ===")
    print({openpyxl.utils.get_column_letter(c): ws.cell(35, c).value for c in range(1, 18)})
    
    print("=== Row 36 ===")
    print({openpyxl.utils.get_column_letter(c): ws.cell(36, c).value for c in range(1, 18)})

if __name__ == "__main__":
    main()
