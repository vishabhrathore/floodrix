import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["24hr SPS"]
    
    print("=== 24hr SPS M31:N37 ===")
    for r in range(31, 38):
        print(f"Row {r}: M={ws.cell(r, 13).value}, N={ws.cell(r, 14).value}")

if __name__ == "__main__":
    main()
