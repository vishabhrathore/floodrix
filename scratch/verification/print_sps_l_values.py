import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["24hr SPS"]
    
    print("=== 24hr SPS Column L and M values ===")
    for r in range(35, 47):
        print(f"Row {r}: L={ws.cell(r, 12).value}, M={ws.cell(r, 13).value}")

if __name__ == "__main__":
    main()
