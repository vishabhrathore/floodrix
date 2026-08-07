import openpyxl

def main():
    path = "/home/vishabh/Downloads/SUH S1_19082025.xlsm"
    wb = openpyxl.load_workbook(path, data_only=True)
    
    uh_24 = [wb["24hr SPS"].cell(r, 2).value for r in range(53, 100)]
    uh_48 = [wb["48hr SPS "].cell(r, 8).value for r in range(35, 82)]
    uh_72 = [wb["72hr SPS  "].cell(r, 8).value for r in range(53, 100)]
    
    print("uh_24 length:", len([x for x in uh_24 if x is not None]))
    print("uh_24 first 10:", uh_24[:10])
    print("uh_48 length:", len([x for x in uh_48 if x is not None]))
    print("uh_48 first 10:", uh_48[:10])
    print("uh_72 length:", len([x for x in uh_72 if x is not None]))
    print("uh_72 first 10:", uh_72[:10])

if __name__ == "__main__":
    main()
